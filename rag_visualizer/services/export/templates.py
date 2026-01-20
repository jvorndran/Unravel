"""Code templates for pipeline export.

Each template is a string with placeholders that get filled in by the generator.
"""

# =============================================================================
# PARSING TEMPLATES
# =============================================================================

PARSING_DOCLING = '''"""Document Parsing with Docling

Advanced PDF parsing with optional OCR and table extraction.

Configuration:
- Enable OCR: {enable_ocr}
- Enable Table Structure: {enable_tables}
- Threads: {num_threads}
"""

import os
import tempfile
from docling.document_converter import DocumentConverter, FormatOption
from docling.datamodel.base_models import InputFormat
from docling.pipeline.standard_pdf_pipeline import StandardPdfPipeline, ThreadedPdfPipelineOptions
from docling.datamodel.accelerator_options import AcceleratorOptions, AcceleratorDevice
from docling.backend.docling_parse_v4_backend import DoclingParseV4DocumentBackend


def parse_pdf(file_path: str) -> str:
    """Parse PDF file using Docling.

    Args:
        file_path: Path to the PDF file

    Returns:
        Extracted text in markdown format
    """
    pipeline_options = ThreadedPdfPipelineOptions(
        accelerator_options=AcceleratorOptions(
            num_threads={num_threads},
            device=AcceleratorDevice.AUTO,
        ),
        do_ocr={enable_ocr},
        do_table_structure={enable_tables},
        generate_page_images=False,
        generate_table_images=False,
        generate_picture_images=False,
    )

    format_option = FormatOption(
        pipeline_options=pipeline_options,
        backend=DoclingParseV4DocumentBackend,
        pipeline_cls=StandardPdfPipeline,
    )

    converter = DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        format_options={{InputFormat.PDF: format_option}},
    )

    result = converter.convert(file_path)
    return result.document.export_to_markdown()


# Example usage:
# text = parse_pdf("document.pdf")
# print(f"Extracted {{len(text)}} characters")
'''

PARSING_DOCX = '''"""Document Parsing for DOCX files

Configuration:
- Output Format: {output_format}
"""

from docx import Document


def parse_docx(file_path: str) -> str:
    """Parse DOCX file and return extracted text.

    Args:
        file_path: Path to the DOCX file

    Returns:
        Extracted text content
    """
    doc = Document(file_path)
    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    return "\\n\\n".join(text_parts)


# Example usage:
# text = parse_docx("document.docx")
# print(f"Extracted {{len(text)}} characters")
'''

PARSING_TEXT = '''"""Document Parsing for Text/Markdown files
"""


def parse_text(file_path: str) -> str:
    """Parse text file and return content.

    Args:
        file_path: Path to the text file

    Returns:
        File content as string
    """
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


# Example usage:
# text = parse_text("document.txt")
# print(f"Read {{len(text)}} characters")
'''

# =============================================================================
# CHUNKING TEMPLATES
# =============================================================================

CHUNKING_HIERARCHICAL = '''"""Text Chunking with Hierarchical Strategy

Structure-aware chunking that creates one chunk per document element (paragraph,
header, list, code block). Best for preserving document structure.

Uses Docling's native HierarchicalChunker for accurate structure detection.

Configuration:
- Merge List Items: {merge_small_chunks}
- Min Chunk Size: {min_chunk_size} characters (for reference, not used by chunker)
"""

import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docling.chunking import HierarchicalChunker
from docling.document_converter import DocumentConverter


@dataclass
class Chunk:
    """A chunk of text with position tracking."""
    text: str
    start_index: int
    end_index: int
    metadata: dict[str, Any]


def _text_to_docling_document(text: str):
    """Convert plain text/markdown to a DoclingDocument for chunking.
    
    Uses Docling's document converter to parse markdown text into a structured document.
    
    Args:
        text: Plain text or markdown content
        
    Returns:
        DoclingDocument instance
    """
    # Create a temporary markdown file
    with tempfile.NamedTemporaryFile(mode='w', suffix='.md', delete=False, encoding='utf-8') as f:
        f.write(text)
        temp_path = f.name
    
    try:
        converter = DocumentConverter()
        result = converter.convert(Path(temp_path))
        return result.document
    finally:
        # Clean up temp file
        try:
            Path(temp_path).unlink()
        except Exception:
            pass


def chunk_text(text: str) -> list[Chunk]:
    """Split text into hierarchical chunks using Docling's HierarchicalChunker.

    Args:
        text: Source text to split

    Returns:
        List of Chunk objects
    """
    merge_list_items = {merge_small_chunks}
    
    # Convert text to DoclingDocument
    try:
        doc = _text_to_docling_document(text)
    except Exception as e:
        # Fallback: return single chunk on conversion error
        return [Chunk(
            text=text,
            start_index=0,
            end_index=len(text),
            metadata={{"strategy": "Hierarchical", "error": str(e), "chunk_index": 0}}
        )]
    
    # Create native HierarchicalChunker
    chunker = HierarchicalChunker(merge_list_items=merge_list_items)
    
    # Generate chunks using native chunker
    try:
        native_chunks = list(chunker.chunk(doc))
    except Exception as e:
        # Fallback: return single chunk on chunking error
        return [Chunk(
            text=text,
            start_index=0,
            end_index=len(text),
            metadata={{"strategy": "Hierarchical", "error": str(e), "chunk_index": 0}}
        )]
    
    # Convert native chunks to our Chunk dataclass
    chunks: list[Chunk] = []
    for i, native_chunk in enumerate(native_chunks):
        chunk_text_str = native_chunk.text
        
        # Extract metadata from native chunk
        metadata: dict[str, Any] = {{
            "strategy": "Hierarchical",
            "chunk_index": i,
            "size": len(chunk_text_str),
        }}
        
        # Extract section hierarchy from headings if available
        if hasattr(native_chunk, "meta") and native_chunk.meta:
            meta = native_chunk.meta
            if hasattr(meta, "headings") and meta.headings:
                metadata["section_hierarchy"] = [
                    h.text if hasattr(h, "text") else str(h) for h in meta.headings
                ]
            
            # Extract element type from doc_items if available
            if hasattr(meta, "doc_items") and meta.doc_items:
                labels = list(set(str(item.label) for item in meta.doc_items))
                metadata["element_type"] = labels if len(labels) > 1 else labels[0] if labels else "text"
        
        # Find chunk position in original text (approximate)
        start_index = text.find(chunk_text_str[:50]) if len(chunk_text_str) >= 50 else text.find(chunk_text_str)
        if start_index == -1:
            start_index = 0
        end_index = start_index + len(chunk_text_str)
        
        chunks.append(Chunk(
            text=chunk_text_str,
            start_index=start_index,
            end_index=end_index,
            metadata=metadata,
        ))
    
    return chunks


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
# for chunk in chunks[:3]:
#     element_type = chunk.metadata.get('element_type', 'text')
#     print(f"[{{element_type}}] {{chunk.text[:100]}}...")
'''

CHUNKING_HYBRID = '''"""Text Chunking with Hybrid Strategy

Token-aware chunking that respects document structure while maintaining token limits.
Best for embedding models with fixed context windows.

Configuration:
- Max Tokens: {max_tokens}
- Chunk Overlap: {chunk_overlap} tokens
- Tokenizer: {tokenizer}
"""

import re
from dataclasses import dataclass
from typing import Any

try:
    import tiktoken
    HAS_TIKTOKEN = True
except ImportError:
    HAS_TIKTOKEN = False


@dataclass
class Chunk:
    """A chunk of text with position tracking."""
    text: str
    start_index: int
    end_index: int
    metadata: dict[str, Any]


def _count_tokens(text: str, tokenizer: str = "{tokenizer}") -> int:
    """Count tokens in text."""
    if HAS_TIKTOKEN:
        try:
            enc = tiktoken.get_encoding(tokenizer)
            return len(enc.encode(text))
        except Exception:
            pass
    return len(text) // 4  # Fallback estimate


def _split_into_elements(text: str) -> list[dict[str, Any]]:
    """Split text into structural elements."""
    elements = []
    current_pos = 0
    header_pattern = r"^(#{{1,6}})\\s+(.+)$"
    list_pattern = r"^(\\s*[-*+]|\\s*\\d+\\.)\\s+(.+)$"

    lines = text.split("\\n")
    i = 0

    while i < len(lines):
        line = lines[i]
        line_start = text.find(line, current_pos)
        if line_start == -1:
            line_start = current_pos

        if line.strip().startswith("```"):
            code_content = [line]
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code_content.append(lines[j])
                j += 1
            if j < len(lines):
                code_content.append(lines[j])
            code_text = "\\n".join(code_content)
            code_end = line_start + len(code_text)
            elements.append({{"type": "code", "text": code_text, "start": line_start, "end": code_end}})
            current_pos = code_end + 1
            i = j + 1
            continue

        header_match = re.match(header_pattern, line)
        if header_match:
            level = len(header_match.group(1))
            elements.append({{"type": f"header_{{level}}", "text": line, "start": line_start, "end": line_start + len(line)}})
            current_pos = line_start + len(line) + 1
            i += 1
            continue

        if re.match(list_pattern, line):
            elements.append({{"type": "list_item", "text": line, "start": line_start, "end": line_start + len(line)}})
            current_pos = line_start + len(line) + 1
            i += 1
            continue

        if not line.strip():
            current_pos = line_start + len(line) + 1
            i += 1
            continue

        para_lines = [line]
        j = i + 1
        while j < len(lines) and lines[j].strip() and not re.match(header_pattern, lines[j]) and not lines[j].strip().startswith("```"):
            if re.match(list_pattern, lines[j]):
                break
            para_lines.append(lines[j])
            j += 1
        para_text = "\\n".join(para_lines)
        para_end = line_start + len(para_text)
        elements.append({{"type": "paragraph", "text": para_text, "start": line_start, "end": para_end}})
        current_pos = para_end + 1
        i = j

    return elements


def _get_overlap_text(text: str, overlap_tokens: int, tokenizer: str) -> str:
    """Get the last N tokens worth of text for overlap."""
    words = text.split()
    overlap_words: list[str] = []
    for word in reversed(words):
        test = " ".join([word] + overlap_words)
        if _count_tokens(test, tokenizer) > overlap_tokens:
            break
        overlap_words.insert(0, word)
    return " ".join(overlap_words)


def chunk_text(text: str) -> list[Chunk]:
    """Split text into token-aware chunks.

    Args:
        text: Source text to split

    Returns:
        List of Chunk objects
    """
    max_tokens = {max_tokens}
    overlap = {chunk_overlap}
    tokenizer = "{tokenizer}"

    elements = _split_into_elements(text)
    if not elements:
        return [Chunk(text=text, start_index=0, end_index=len(text),
                      metadata={{"strategy": "Hybrid", "chunk_index": 0, "size": len(text)}})]

    chunks = []
    current_text = ""
    current_start = 0
    current_headers: list[str] = []

    def make_chunk(txt: str, start: int, idx: int, headers: list[str]) -> Chunk:
        metadata = {{"strategy": "Hybrid", "chunk_index": idx, "size": len(txt)}}
        if headers:
            metadata["section_hierarchy"] = headers
        return Chunk(text=txt, start_index=start, end_index=start + len(txt), metadata=metadata)

    for elem in elements:
        if elem["type"].startswith("header_"):
            level = int(elem["type"].split("_")[1])
            current_headers = current_headers[:level - 1]
            current_headers.append(elem["text"].lstrip("#").strip())

        elem_text = elem["text"]
        elem_tokens = _count_tokens(elem_text, tokenizer)

        if elem_tokens > max_tokens:
            if current_text.strip():
                chunks.append(make_chunk(current_text, current_start, len(chunks), current_headers.copy()))
            # Split large element by words
            words = elem_text.split()
            chunk_words: list[str] = []
            chunk_start = elem["start"]
            for word in words:
                test_text = " ".join(chunk_words + [word])
                if _count_tokens(test_text, tokenizer) > max_tokens and chunk_words:
                    chunk_text_str = " ".join(chunk_words)
                    chunks.append(make_chunk(chunk_text_str, chunk_start, len(chunks), current_headers.copy()))
                    if overlap > 0:
                        overlap_words = []
                        for w in reversed(chunk_words):
                            test = " ".join([w] + overlap_words)
                            if _count_tokens(test, tokenizer) > overlap:
                                break
                            overlap_words.insert(0, w)
                        chunk_words = overlap_words + [word]
                    else:
                        chunk_words = [word]
                    chunk_start = elem["start"] + elem_text.find(" ".join(chunk_words))
                else:
                    chunk_words.append(word)
            if chunk_words:
                chunk_text_str = " ".join(chunk_words)
                chunks.append(make_chunk(chunk_text_str, chunk_start, len(chunks), current_headers.copy()))
            current_text = ""
            current_start = elem["end"] + 1
            continue

        combined = current_text + ("\\n\\n" if current_text else "") + elem_text
        if _count_tokens(combined, tokenizer) > max_tokens and current_text.strip():
            chunks.append(make_chunk(current_text, current_start, len(chunks), current_headers.copy()))
            if overlap > 0:
                overlap_text = _get_overlap_text(current_text, overlap, tokenizer)
                current_text = overlap_text + "\\n\\n" + elem_text if overlap_text else elem_text
            else:
                current_text = elem_text
            current_start = elem["start"]
        else:
            if not current_text:
                current_start = elem["start"]
            current_text = combined

    if current_text.strip():
        chunks.append(make_chunk(current_text, current_start, len(chunks), current_headers.copy()))

    return chunks


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
# for chunk in chunks[:3]:
#     print(f"[{{_count_tokens(chunk.text)}} tokens] {{chunk.text[:100]}}...")
'''

CHUNKING_RECURSIVE = '''"""Text Chunking with RecursiveCharacterTextSplitter

Best for most documents. Splits by paragraphs, then sentences, then words
to maintain semantic coherence.

Configuration:
- Chunk Size: {chunk_size} characters
- Chunk Overlap: {chunk_overlap} characters
"""

from langchain_text_splitters import RecursiveCharacterTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split text into overlapping chunks.

    Args:
        text: Source text to split

    Returns:
        List of text chunks
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
        length_function=len,
        add_start_index=True,
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
# for i, chunk in enumerate(chunks[:3]):
#     print(f"Chunk {{i+1}}: {{chunk[:100]}}...")
'''

CHUNKING_CHARACTER = '''"""Text Chunking with CharacterTextSplitter

Simple split on a separator. Use for well-structured text with consistent delimiters.

Configuration:
- Chunk Size: {chunk_size} characters
- Chunk Overlap: {chunk_overlap} characters
- Separator: {separator_display}
"""

from langchain_text_splitters import CharacterTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split text into chunks on separator.

    Args:
        text: Source text to split

    Returns:
        List of text chunks
    """
    splitter = CharacterTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
        separator={separator},
        length_function=len,
        add_start_index=True,
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_TOKEN = '''"""Text Chunking with TokenTextSplitter

Splits by tokens (not characters). Use when you need precise token counts
for LLM context windows.

Configuration:
- Chunk Size: {chunk_size} tokens
- Chunk Overlap: {chunk_overlap} tokens
- Encoding: {encoding_name}
"""

from langchain_text_splitters import TokenTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split text into chunks by token count.

    Args:
        text: Source text to split

    Returns:
        List of text chunks
    """
    splitter = TokenTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
        encoding_name="{encoding_name}",
        add_start_index=True,
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_MARKDOWN = '''"""Text Chunking with MarkdownTextSplitter

Use for Markdown files. Keeps headings with their content for better semantic chunks.

Configuration:
- Chunk Size: {chunk_size} characters
- Chunk Overlap: {chunk_overlap} characters
"""

from langchain_text_splitters import MarkdownTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split markdown text into semantic chunks.

    Args:
        text: Markdown text to split

    Returns:
        List of text chunks
    """
    splitter = MarkdownTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
        add_start_index=True,
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(markdown_text)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_LATEX = '''"""Text Chunking with LatexTextSplitter

Use for LaTeX/academic documents. Respects sections, equations, and environments.

Configuration:
- Chunk Size: {chunk_size} characters
- Chunk Overlap: {chunk_overlap} characters
"""

from langchain_text_splitters import LatexTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split LaTeX text into semantic chunks.

    Args:
        text: LaTeX text to split

    Returns:
        List of text chunks
    """
    splitter = LatexTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
        add_start_index=True,
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(latex_text)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_PYTHON = '''"""Text Chunking with PythonCodeTextSplitter

Use for Python code. Splits at function/class boundaries to keep code units intact.

Configuration:
- Chunk Size: {chunk_size} characters
- Chunk Overlap: {chunk_overlap} characters
"""

from langchain_text_splitters import PythonCodeTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split Python code into semantic chunks.

    Args:
        text: Python code to split

    Returns:
        List of code chunks
    """
    splitter = PythonCodeTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
        add_start_index=True,
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(python_code)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_HTML = '''"""Text Chunking with HTMLHeaderTextSplitter

Use for HTML pages. Creates chunks based on heading hierarchy (h1, h2, h3).
"""

from langchain_text_splitters import HTMLHeaderTextSplitter


def chunk_text(html: str) -> list[str]:
    """Split HTML into chunks based on headers.

    Args:
        html: HTML content to split

    Returns:
        List of text chunks
    """
    headers_to_split_on = [
        ("h1", "Header 1"),
        ("h2", "Header 2"),
        ("h3", "Header 3"),
    ]

    splitter = HTMLHeaderTextSplitter(headers_to_split_on=headers_to_split_on)
    documents = splitter.split_text(html)

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(html_content)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_JSON = '''"""Text Chunking with RecursiveJsonSplitter

Use for JSON data. Keeps nested objects together while respecting size limits.

Configuration:
- Max Chunk Size: {max_chunk_size} characters
- Min Chunk Size: {min_chunk_size} characters
"""

import json
from langchain_text_splitters import RecursiveJsonSplitter


def chunk_json(json_text: str) -> list[str]:
    """Split JSON into semantic chunks.

    Args:
        json_text: JSON string to split

    Returns:
        List of JSON string chunks
    """
    splitter = RecursiveJsonSplitter(
        max_chunk_size={max_chunk_size},
        min_chunk_size={min_chunk_size},
    )

    json_data = json.loads(json_text)
    json_chunks = splitter.split_json(json_data)

    return [json.dumps(chunk, indent=2) for chunk in json_chunks]


# Example usage:
# chunks = chunk_json(json_string)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_SENTENCE_TRANSFORMERS = '''"""Text Chunking with SentenceTransformersTokenTextSplitter

Aligns with embedding model tokenization. Use for optimal embedding boundaries.

Configuration:
- Tokens Per Chunk: {tokens_per_chunk}
- Chunk Overlap: {chunk_overlap} tokens
- Model: {model_name}
"""

from langchain_text_splitters import SentenceTransformersTokenTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split text aligned to embedding model tokens.

    Args:
        text: Text to split

    Returns:
        List of text chunks
    """
    splitter = SentenceTransformersTokenTextSplitter(
        tokens_per_chunk={tokens_per_chunk},
        chunk_overlap={chunk_overlap},
        model_name="{model_name}",
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_NLTK = '''"""Text Chunking with NLTKTextSplitter

Uses NLTK for sentence detection. Good for natural language with proper punctuation.

Configuration:
- Chunk Size: {chunk_size} characters
- Chunk Overlap: {chunk_overlap} characters
"""

from langchain_text_splitters import NLTKTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split text using NLTK sentence boundaries.

    Args:
        text: Text to split

    Returns:
        List of text chunks
    """
    splitter = NLTKTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
'''

CHUNKING_SPACY = '''"""Text Chunking with SpacyTextSplitter

Uses spaCy NLP for intelligent text segmentation. Best for complex natural language.

Configuration:
- Chunk Size: {chunk_size} characters
- Chunk Overlap: {chunk_overlap} characters
- Pipeline: {pipeline}
"""

from langchain_text_splitters import SpacyTextSplitter


def chunk_text(text: str) -> list[str]:
    """Split text using spaCy NLP pipeline.

    Args:
        text: Text to split

    Returns:
        List of text chunks
    """
    splitter = SpacyTextSplitter(
        chunk_size={chunk_size},
        chunk_overlap={chunk_overlap},
        pipeline="{pipeline}",
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
'''

# Generic fallback
CHUNKING_GENERIC = '''"""Text Chunking with {splitter_name}

Configuration:
{config_comment}
"""

from langchain_text_splitters import {splitter_name}


def chunk_text(text: str) -> list[str]:
    """Split text into chunks.

    Args:
        text: Text to split

    Returns:
        List of text chunks
    """
    splitter = {splitter_name}(
{params_code}
    )

    documents = splitter.create_documents([text])

    return [doc.page_content for doc in documents]


# Example usage:
# chunks = chunk_text(text)
# print(f"Created {{len(chunks)}} chunks")
'''

# =============================================================================
# EMBEDDING TEMPLATE
# =============================================================================

EMBEDDING_TEMPLATE = '''"""Embedding Generation with SentenceTransformers

Model: {model_name}
Dimension: {dimension}
Description: {description}
"""

from sentence_transformers import SentenceTransformer
import numpy as np

# Initialize the embedding model
model = SentenceTransformer("{model_name}")


def generate_embeddings(texts: list[str], batch_size: int = 32) -> np.ndarray:
    """Generate embeddings for a list of texts.

    Args:
        texts: List of text strings to embed
        batch_size: Number of texts to process at once

    Returns:
        numpy array of shape (len(texts), {dimension})
    """
    embeddings = model.encode(
        texts,
        batch_size=batch_size,
        show_progress_bar=True,
        normalize_embeddings=True,  # Recommended for cosine similarity
        convert_to_numpy=True,
    )

    return embeddings


def embed_query(query: str) -> np.ndarray:
    """Generate embedding for a single query.

    Args:
        query: Query text to embed

    Returns:
        numpy array of shape ({dimension},)
    """
    return model.encode(query, normalize_embeddings=True, convert_to_numpy=True)


# Example usage:
# texts = ["First chunk", "Second chunk", "Third chunk"]
# embeddings = generate_embeddings(texts)
# print(f"Generated embeddings with shape: {{embeddings.shape}}")
#
# # Similarity search example
# query_embedding = embed_query("search query")
# similarities = np.dot(embeddings, query_embedding)  # Cosine similarity (normalized)
# print(f"Similarities: {{similarities}}")
'''

# =============================================================================
# TEMPLATE REGISTRIES
# =============================================================================

PARSING_TEMPLATES = {
    "docling": PARSING_DOCLING,
}

CHUNKING_TEMPLATES = {
    # Docling chunkers (primary)
    "HierarchicalChunker": CHUNKING_HIERARCHICAL,
    "HybridChunker": CHUNKING_HYBRID,
    # Legacy LangChain splitters (for backwards compatibility in exports)
    "RecursiveCharacterTextSplitter": CHUNKING_RECURSIVE,
    "CharacterTextSplitter": CHUNKING_CHARACTER,
    "TokenTextSplitter": CHUNKING_TOKEN,
    "MarkdownTextSplitter": CHUNKING_MARKDOWN,
    "LatexTextSplitter": CHUNKING_LATEX,
    "PythonCodeTextSplitter": CHUNKING_PYTHON,
    "HTMLHeaderTextSplitter": CHUNKING_HTML,
    "RecursiveJsonSplitter": CHUNKING_JSON,
    "SentenceTransformersTokenTextSplitter": CHUNKING_SENTENCE_TRANSFORMERS,
    "NLTKTextSplitter": CHUNKING_NLTK,
    "SpacyTextSplitter": CHUNKING_SPACY,
}
